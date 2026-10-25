#!/usr/bin/env python3
"""Overlay Yuni/Prisha scores per Excel tab. Requires openpyxl, numpy, matplotlib."""
import argparse
import csv
import math
import re
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

LABELS = ['[1, 2)', '[2, 3)', '[3, 4)', '[4, 5)', '5']

def counts(values):
    a = np.asarray(values)
    return [int(np.sum((a >= i) & (a < i + 1))) for i in range(1, 5)] + [int(np.sum(a == 5))]

def draw(ax, title, scores):
    for name, color in [('Yuni', 'tab:blue'), ('Prisha', 'tab:orange')]:
        ax.stairs(counts(scores[name]), np.arange(6) - .5, fill=True,
                  color=color, alpha=.38, label=f'{name} (n={len(scores[name])})')
    ax.set(title=title.replace('_', ' '), xlabel='Visual-inspection score', ylabel='Number of targets', xlim=(-.5, 4.5))
    ax.set_xticks(range(5), LABELS)
    ax.set_ylim(bottom=0)
    ax.spines[['top', 'right']].set_visible(False)
    ax.grid(axis='y', alpha=.18)
    ax.set_axisbelow(True)
    ax.legend(frameon=False, fontsize=9)

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--out-dir', type=Path, default=Path('histograms'))
    args = parser.parse_args()
    args.out_dir.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    categories = []
    for sheet in workbook:
        rows = sheet.iter_rows(values_only=True)
        header = [str(x or '').strip().lower() for x in next(rows, ())]
        if 'targetid' not in header:
            continue
        columns = {}
        for name in ('Yuni', 'Prisha'):
            matches = [i for i, h in enumerate(header) if name.lower() in h and 'score' in h]
            if matches:
                columns[name] = matches[0]
        if len(columns) != 2:
            print(f'Skipping {sheet.title}: missing score columns')
            continue
        scores = {name: [] for name in columns}
        for row in rows:
            # Ignore blank rows and trailing summary means without a TARGETID.
            if row[header.index('targetid')] in (None, ''):
                continue
            for name, column in columns.items():
                value = row[column]
                if value in (None, ''):
                    continue
                try:
                    value = float(value)
                except (TypeError, ValueError):
                    print(f'Ignoring nonnumeric {name} score in {sheet.title}: {value!r}')
                    continue
                if math.isfinite(value) and 1 <= value <= 5:
                    scores[name].append(value)
                else:
                    print(f'Ignoring out-of-range {name} score in {sheet.title}: {value}')
        categories.append((sheet.title, scores))
        fig, ax = plt.subplots(figsize=(8, 5), layout='constrained')
        draw(ax, sheet.title, scores)
        slug = re.sub(r'[^a-zA-Z0-9_-]+', '_', sheet.title).strip('_')
        fig.savefig(args.out_dir / f'{slug}.png', dpi=200)
        plt.close(fig)
        print(sheet.title, {n: counts(v) for n, v in scores.items()})
    workbook.close()
    if not categories:
        raise ValueError('No tabs with TARGETID and both score columns found')
    fig, axes = plt.subplots(math.ceil(len(categories)/2), 2, figsize=(13, 3.8*math.ceil(len(categories)/2)), layout='constrained')
    for ax, (title, scores) in zip(axes.flat, categories):
        draw(ax, title, scores)
    for ax in list(axes.flat)[len(categories):]:
        ax.set_visible(False)
    fig.suptitle('Visual-inspection scores', fontsize=18)
    fig.savefig(args.out_dir / 'all_categories.png', dpi=180)
    plt.close(fig)
    with (args.out_dir / 'bin_counts.csv').open('w', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['category', 'reviewer', 'n', *LABELS])
        for title, scores in categories:
            for name, values in scores.items():
                writer.writerow([title, name, len(values), *counts(values)])

if __name__ == '__main__':
    main()
